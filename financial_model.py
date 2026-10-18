"""Historical modeling inputs and portable Excel export for the deployed app."""
from io import BytesIO
import math
import pandas as pd

SECTIONS = {
    '손익계산서': ['매출액','매출원가','매출총이익','판매관리비','영업이익','법인세차감전이익','법인세비용','순이익'],
    '재무상태표': ['유동자산','현금및현금성자산','재고자산','유형자산','자산','유동부채','부채','자본'],
    '현금흐름표': ['영업현금흐름','투자현금흐름','재무현금흐름','유형자산취득'],
}
RATIOS = ['매출 증가율','매출총이익률','영업이익률','순이익률','유동비율','부채비율','ROE','FCF (유형자산 기준)']

def numeric(v):
    try:
        v=float(v)
        return v if math.isfinite(v) else None
    except (ValueError,TypeError):return None

def calculated(cur, prev):
    def divide(a,b):
        x,y=numeric(cur.get(a)),numeric(cur.get(b))
        return x/y if x is not None and y is not None and y>0 else None
    revenue,old=numeric(cur.get('매출액')),numeric(prev.get('매출액'))
    equity,prior,profit=[numeric(x) for x in (cur.get('자본'),prev.get('자본'),cur.get('순이익'))]
    cash,capex=numeric(cur.get('영업현금흐름')),numeric(cur.get('유형자산취득'))
    return dict(zip(RATIOS,[
        revenue/old-1 if revenue is not None and old is not None and old>0 else None,
        divide('매출총이익','매출액'),divide('영업이익','매출액'),divide('순이익','매출액'),
        divide('유동자산','유동부채'),divide('부채','자본'),
        profit/((equity+prior)/2) if None not in (profit,equity,prior) and equity>0 and prior>0 else None,
        cash-abs(capex) if cash is not None and capex is not None else None]))

def grid(rows, labels, years, derived=False):
    indexed={int(r['연도']):r for r in rows}
    return pd.DataFrame({str(y):[numeric((calculated(indexed.get(y,{}),indexed.get(y-1,{})) if derived else indexed.get(y,{})).get(k)) for k in labels] for y in years},index=pd.Index(labels,name='항목'))

def export_xlsx(saved, years, basis):
    # Runtime export must work on Streamlit Cloud without a Node authoring runtime.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook();wb.remove(wb.active)
    rows=saved['rows']; by_year={int(r['연도']):r for r in rows}
    title=saved['company']['name']+' · '+('연결' if basis=='CFS' else '별도')+' · 연간 실적 · 억 원'
    mapping={}
    for section,labels in SECTIONS.items():
        ws=wb.create_sheet(section);ws.append([title]);ws.append(['항목']+list(years))
        for label in labels:
            ws.append([label]+[None if numeric(by_year.get(y,{}).get(label)) is None else numeric(by_year[y][label])/1e8 for y in years])
            mapping[label]=(section,ws.max_row)
        ws.append(['출처: 각 연도 사업보고서의 당기금액. 미제공·모호한 계정은 공란.'])
    ratios=wb.create_sheet('핵심 지표',0);ratios.append([title]);ratios.append(['항목']+list(years))
    for label in RATIOS:ratios.append([label])
    def ref(label,col):
        sheet,row=mapping[label];return "'%s'!%s%d"%(sheet,get_column_letter(col),row)
    for col,y in enumerate(years,2):
        def ratio(a,b):
            x,z=ref(a,col),ref(b,col)
            return f'=IF(AND(ISNUMBER({x}),ISNUMBER({z}),{z}>0),{x}/{z},"")'
        rev=ref('매출액',col)
        if y-1 in years:
            old=ref('매출액',years.index(y-1)+2)
            growth=f'=IF(AND(ISNUMBER({rev}),ISNUMBER({old}),{old}>0),{rev}/{old}-1,"")'
        else:
            # Extra prior year fetched for first visible year's growth and ROE.
            growth=calculated(by_year.get(y,{}),by_year.get(y-1,{}))['매출 증가율']
        equity=ref('자본',col);profit=ref('순이익',col)
        if y-1 in years:
            oldeq=ref('자본',years.index(y-1)+2)
            roe=f'=IF(AND(ISNUMBER({profit}),ISNUMBER({equity}),ISNUMBER({oldeq}),{equity}>0,{oldeq}>0),{profit}/AVERAGE({equity},{oldeq}),"")'
        else:roe=calculated(by_year.get(y,{}),by_year.get(y-1,{}))['ROE']
        cash,capex=ref('영업현금흐름',col),ref('유형자산취득',col)
        values=[growth,ratio('매출총이익','매출액'),ratio('영업이익','매출액'),ratio('순이익','매출액'),ratio('유동자산','유동부채'),ratio('부채','자본'),roe,f'=IF(AND(ISNUMBER({cash}),ISNUMBER({capex})),{cash}-ABS({capex}),"")']
        for row,value in enumerate(values,3):ratios.cell(row,col,value)
    ratios.append(['FCF = 영업현금흐름 − 유형자산취득 절대값. 무형자산 취득은 제외.'])
    ratios.append(['ROE = 총순이익 / 기초·기말 평균 총자본. 분모 0 이하이면 공란.'])
    ratios.append(['첫 표시연도 성장률·ROE는 추가 조회한 전년도 기준 계산값. 예측·정상화 조정 미포함.'])
    raw=wb.create_sheet('DART 원자료')
    raw.append(['사업연도','재무제표','계정 ID','계정명','통화','당기금액 (원)','접수번호','원문'])
    for y,records in sorted(saved['raw'].items()):
        for r in records:
            receipt=r.get('rcept_no','')
            raw.append([y,r.get('sj_div'),r.get('account_id'),r.get('account_nm'),r.get('currency'),numeric(str(r.get('thstrm_amount','')).replace(',','')),receipt,'https://dart.fss.or.kr/dsaf001/main.do?rcpNo='+receipt])
    for ws in wb:
        ws.freeze_panes='B3' if ws!=raw else 'F2';ws.sheet_view.showGridLines=False
        ws.column_dimensions['A'].width=30
        for col in range(2,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=19 if ws!=raw else 25
        for row in ws:
            for cell in row:
                cell.font=Font(name='맑은 고딕',size=10,color='152D43')
                if cell.row==(1 if ws==raw else 2):
                    cell.fill=PatternFill('solid',fgColor='152D43');cell.font=Font(name='맑은 고딕',size=10,bold=True,color='FFFFFF')
                elif cell.row>2 and cell.column>1 and ws!=raw:
                    cell.number_format='0.0%;[Red](0.0%);"—"' if ws==ratios and cell.row<10 else '#,##0.0;[Red](#,##0.0);"—"'
                    cell.alignment=Alignment(horizontal='right')
                # Treat upstream names as literal text, never spreadsheet formulas.
                if isinstance(cell.value,str) and cell.value.startswith(('=','+','-','@')) and not (ws==ratios and cell.column>1 and cell.value.startswith('=')):
                    cell.data_type='s'
        ws.sheet_properties.pageSetUpPr.fitToPage=True
    wb.calculation.fullCalcOnLoad=True
    out=BytesIO();wb.save(out);return out.getvalue()
